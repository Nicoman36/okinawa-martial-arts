#!/usr/bin/env python3
"""Export data/locations.csv to an Excel workbook (for sharing / spreadsheet editing)."""
import csv, sys
from collections import Counter
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
sys.path.insert(0, str(Path(__file__).resolve().parent))
from catalog import DISC, VERIF
ROOT = Path(__file__).resolve().parent.parent
rows = list(csv.DictReader(open(ROOT/"data/locations.csv", encoding="utf-8-sig")))
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Directory"
hdr = ["Art(s)","Name (English)","Name (日本語)","Style","Type","Area","Address (日本語)","Where (English)","Phone","Website","Social","Email","Hours / classes","Contact / instructor","Visitors / trial","English?","Verification","Status","Last verified","Notes","ID"]
ws.append(hdr)
for r in sorted(rows, key=lambda r: (r["disciplines"].split(";")[0], r["municipality"], r["name_en"])):
    ws.append([", ".join(DISC[d]["en"] for d in r["disciplines"].split(";")), r["name_en"], r["name_ja"], r["style"], r["kind"], r["municipality"], r["address_ja"], r["address_en"], r["phone"], r["website"], r["social"], r["email"], r["hours"], r["contact"], r["visitors"], r["english"], f'{r["verification"]} — {VERIF[r["verification"]][0]}', r["status"], r["last_verified"], r["notes"], r["id"]])
for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F2F5A"); c.alignment = Alignment(wrap_text=True, vertical="center")
for col, w in zip("ABCDEFGHIJKLMNOPQRSTU", [26,38,34,28,10,14,34,30,15,30,26,24,34,30,34,14,30,12,12,40,22]): ws.column_dimensions[col].width = w
ws.freeze_panes = "C2"; ws.auto_filter.ref = ws.dimensions
nonorg = [r for r in rows if r["kind"] != "organization"]
s = wb.create_sheet("By art"); s.append(["Art","Places to train","Checked (A)","Listed (B)","Roster 2019 (C)","Lead (D)"])
for d in DISC.values():
    rs = [r for r in nonorg if d["slug"] in r["disciplines"].split(";")]
    c = Counter(r["verification"] for r in rs); s.append([d["en"], len(rs), c["A"], c["B"], c["C"], c["D"]])
a = wb.create_sheet("By area"); a.append(["Municipality","Places to train"])
for k, v in Counter(r["municipality"] or "(not listed)" for r in nonorg).most_common(): a.append([k, v])
for sh in (s, a):
    for c in sh[1]: c.font = Font(bold=True)
    sh.column_dimensions["A"].width = 34
out = ROOT/"dist_extras"; out.mkdir(exist_ok=True)
wb.save(out/"Okinawa_Martial_Arts_Directory_v2.xlsx"); print("saved")
